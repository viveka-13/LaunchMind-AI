"""
agent/financial_model_generator.py
Generates Financial_Model.xlsx with 3 sheets:
  1. Pricing Scenarios (12-month MRR simulation)
  2. 3-Year P&L (36-month projection with chart)
  3. Cap Table (funding & equity)

ALL calculated values use Excel formulas. No hardcoded computed values.
"""

import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.chart import LineChart, Reference
from openpyxl.utils import get_column_letter


# ─── Shared Styles ───────────────────────────

HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
INPUT_FILL = PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid")  # Yellow = editable
LABEL_FONT = Font(name="Calibri", size=10, bold=True, color="1F4E79")
NUM_FONT = Font(name="Calibri", size=10)
CURRENCY_FMT = '#,##0'
PCT_FMT = '0.0%'
THIN_BORDER = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9'),
)


def _style_header(ws, row, max_col):
    for c in range(1, max_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', wrap_text=True)
        cell.border = THIN_BORDER


def _style_cell(cell, fmt=None, is_input=False, is_label=False):
    cell.border = THIN_BORDER
    cell.alignment = Alignment(horizontal='right' if not is_label else 'left')
    if fmt:
        cell.number_format = fmt
    if is_input:
        cell.fill = INPUT_FILL
    if is_label:
        cell.font = LABEL_FONT
    else:
        cell.font = NUM_FONT


# ═══════════════════════════════════════════════
# Sheet 1: Pricing Scenarios
# ═══════════════════════════════════════════════

def build_pricing_sheet(wb: Workbook, pricing_data: dict):
    ws = wb.create_sheet("Pricing Scenarios", 0)
    ws.sheet_properties.tabColor = "3B82F6"

    tiers = pricing_data.get("tiers", [])
    if not tiers:
        tiers = [
            {"name": "Basic", "monthly_price": 29, "target_segment": "Small biz", "adoption_pct": 0.05, "monthly_growth_rate": 0.08},
            {"name": "Pro", "monthly_price": 79, "target_segment": "Medium biz", "adoption_pct": 0.03, "monthly_growth_rate": 0.06},
        ]
    total_customers = pricing_data.get("total_addressable_customers", 10000)
    currency = pricing_data.get("currency", "USD")

    # ── Section A: Assumptions ──
    ws.merge_cells('A1:E1')
    ws['A1'] = "PRICING MODEL ASSUMPTIONS"
    ws['A1'].font = Font(name="Calibri", size=13, bold=True, color="1F4E79")

    ws['A3'] = "Total Addressable Customers"
    _style_cell(ws['A3'], is_label=True)
    ws['B3'] = total_customers
    _style_cell(ws['B3'], fmt=CURRENCY_FMT, is_input=True)

    # Tier headers
    row = 5
    headers = ["Tier Name", f"Price ({currency}/mo)", "Target Segment", "Adoption %", "Monthly Growth Rate"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=row, column=i, value=h)
    _style_header(ws, row, len(headers))

    tier_start_row = row + 1
    for idx, tier in enumerate(tiers):
        r = tier_start_row + idx
        ws.cell(row=r, column=1, value=tier.get("name", f"Tier {idx+1}"))
        _style_cell(ws.cell(row=r, column=1), is_label=True)

        ws.cell(row=r, column=2, value=tier.get("monthly_price", 29))
        _style_cell(ws.cell(row=r, column=2), fmt=CURRENCY_FMT, is_input=True)

        ws.cell(row=r, column=3, value=tier.get("target_segment", ""))
        _style_cell(ws.cell(row=r, column=3), is_input=True)

        ws.cell(row=r, column=4, value=tier.get("adoption_pct", 0.05))
        _style_cell(ws.cell(row=r, column=4), fmt=PCT_FMT, is_input=True)

        ws.cell(row=r, column=5, value=tier.get("monthly_growth_rate", 0.06))
        _style_cell(ws.cell(row=r, column=5), fmt=PCT_FMT, is_input=True)

    # ── Section B: 12-Month MRR Simulation ──
    sim_start_row = tier_start_row + len(tiers) + 2
    ws.cell(row=sim_start_row, column=1, value="12-MONTH MRR SIMULATION")
    ws.cell(row=sim_start_row, column=1).font = Font(name="Calibri", size=13, bold=True, color="1F4E79")

    header_row = sim_start_row + 1
    ws.cell(row=header_row, column=1, value="Metric")
    for m in range(1, 13):
        ws.cell(row=header_row, column=m + 1, value=f"Month {m}")
    _style_header(ws, header_row, 13)

    # For each tier: Customers row, then MRR row
    current_row = header_row + 1
    mrr_rows = []

    for idx, tier in enumerate(tiers):
        tier_data_row = tier_start_row + idx  # row where tier price/adoption is
        tier_name = tier.get("name", f"Tier {idx+1}")

        # Customers row
        cust_row = current_row
        ws.cell(row=cust_row, column=1, value=f"{tier_name} - Customers")
        _style_cell(ws.cell(row=cust_row, column=1), is_label=True)

        # B3 = total customers, D{tier_data_row} = adoption %, E{tier_data_row} = growth
        adoption_cell = f"$D${tier_data_row}"
        growth_cell = f"$E${tier_data_row}"
        total_cust_cell = "$B$3"

        for m in range(1, 13):
            col = m + 1
            col_letter = get_column_letter(col)
            if m == 1:
                formula = f"=ROUND({total_cust_cell}*{adoption_cell},0)"
            else:
                prev_col = get_column_letter(col - 1)
                formula = f"=ROUND({prev_col}{cust_row}*(1+{growth_cell}),0)"
            ws.cell(row=cust_row, column=col, value=formula)
            _style_cell(ws.cell(row=cust_row, column=col), fmt=CURRENCY_FMT)

        # MRR row
        mrr_row = current_row + 1
        mrr_rows.append(mrr_row)
        ws.cell(row=mrr_row, column=1, value=f"{tier_name} - MRR")
        _style_cell(ws.cell(row=mrr_row, column=1), is_label=True)

        price_cell = f"$B${tier_data_row}"
        for m in range(1, 13):
            col = m + 1
            col_letter = get_column_letter(col)
            formula = f"={col_letter}{cust_row}*{price_cell}"
            ws.cell(row=mrr_row, column=col, value=formula)
            _style_cell(ws.cell(row=mrr_row, column=col), fmt=CURRENCY_FMT)

        current_row += 2

    # Total MRR row
    total_row = current_row
    ws.cell(row=total_row, column=1, value="TOTAL MRR")
    ws.cell(row=total_row, column=1).font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    ws.cell(row=total_row, column=1).fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")

    for m in range(1, 13):
        col = m + 1
        col_letter = get_column_letter(col)
        refs = "+".join([f"{col_letter}{r}" for r in mrr_rows])
        formula = f"={refs}"
        ws.cell(row=total_row, column=col, value=formula)
        cell = ws.cell(row=total_row, column=col)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
        cell.number_format = CURRENCY_FMT
        cell.border = THIN_BORDER

    # Column widths
    ws.column_dimensions['A'].width = 28
    for c in range(2, 14):
        ws.column_dimensions[get_column_letter(c)].width = 14

    return total_row  # return for cross-sheet reference


# ═══════════════════════════════════════════════
# Sheet 2: 3-Year P&L
# ═══════════════════════════════════════════════

def build_projection_sheet(wb: Workbook, projection_data: dict, pricing_total_mrr_row: int, num_tiers: int):
    ws = wb.create_sheet("3-Year P&L", 1)
    ws.sheet_properties.tabColor = "8B5CF6"

    cogs_pct = projection_data.get("cogs_pct", 0.25)
    expenses = projection_data.get("operating_expenses", {})
    expense_growth = projection_data.get("expense_growth_rate", 0.05)
    hires = projection_data.get("hires", [])
    initial_cash = projection_data.get("initial_cash", 100000)
    currency = projection_data.get("currency", "USD")

    # ── Assumptions section ──
    ws.merge_cells('A1:D1')
    ws['A1'] = "3-YEAR P&L ASSUMPTIONS"
    ws['A1'].font = Font(name="Calibri", size=13, bold=True, color="1F4E79")

    ws['A3'] = "COGS %"
    _style_cell(ws['A3'], is_label=True)
    ws['B3'] = cogs_pct
    _style_cell(ws['B3'], fmt=PCT_FMT, is_input=True)

    ws['A4'] = "Quarterly Expense Growth"
    _style_cell(ws['A4'], is_label=True)
    ws['B4'] = expense_growth
    _style_cell(ws['B4'], fmt=PCT_FMT, is_input=True)

    ws['A5'] = f"Initial Cash ({currency})"
    _style_cell(ws['A5'], is_label=True)
    ws['B5'] = initial_cash
    _style_cell(ws['B5'], fmt=CURRENCY_FMT, is_input=True)

    # Operating expense assumptions
    ws['A7'] = "MONTHLY OPERATING EXPENSES"
    ws['A7'].font = Font(name="Calibri", size=11, bold=True, color="1F4E79")

    exp_start = 8
    exp_labels = list(expenses.keys())
    for i, (label, val) in enumerate(expenses.items()):
        r = exp_start + i
        ws.cell(row=r, column=1, value=label.replace("_", " ").title())
        _style_cell(ws.cell(row=r, column=1), is_label=True)
        ws.cell(row=r, column=2, value=val)
        _style_cell(ws.cell(row=r, column=2), fmt=CURRENCY_FMT, is_input=True)

    # Hires assumptions
    hire_header_row = exp_start + len(exp_labels) + 1
    ws.cell(row=hire_header_row, column=1, value="HIRING PLAN")
    ws.cell(row=hire_header_row, column=1).font = Font(name="Calibri", size=11, bold=True, color="1F4E79")

    hire_cols = ["Role", f"Salary ({currency}/mo)", "Start Month"]
    hire_data_start = hire_header_row + 1
    for i, h in enumerate(hire_cols, 1):
        ws.cell(row=hire_data_start, column=i, value=h)
    _style_header(ws, hire_data_start, len(hire_cols))

    for i, hire in enumerate(hires):
        r = hire_data_start + 1 + i
        ws.cell(row=r, column=1, value=hire.get("role", f"Hire {i+1}"))
        _style_cell(ws.cell(row=r, column=1), is_label=True)
        ws.cell(row=r, column=2, value=hire.get("monthly_salary", 5000))
        _style_cell(ws.cell(row=r, column=2), fmt=CURRENCY_FMT, is_input=True)
        ws.cell(row=r, column=3, value=hire.get("start_month", 1))
        _style_cell(ws.cell(row=r, column=3), is_input=True)

    # ── 36-Month P&L Table ──
    table_start = hire_data_start + 1 + len(hires) + 2
    ws.cell(row=table_start, column=1, value="36-MONTH PROFIT & LOSS")
    ws.cell(row=table_start, column=1).font = Font(name="Calibri", size=13, bold=True, color="1F4E79")

    header_row = table_start + 1
    ws.cell(row=header_row, column=1, value="")
    for m in range(1, 37):
        ws.cell(row=header_row, column=m + 1, value=f"M{m}")
    _style_header(ws, header_row, 37)

    # Row definitions
    rev_row = header_row + 1
    cogs_row = rev_row + 1
    gross_row = cogs_row + 1
    opex_row = gross_row + 1
    hire_cost_row = opex_row + 1
    total_exp_row = hire_cost_row + 1
    profit_row = total_exp_row + 1
    cash_row = profit_row + 1

    row_labels = [
        (rev_row, "Revenue"),
        (cogs_row, "COGS"),
        (gross_row, "Gross Margin"),
        (opex_row, "Operating Expenses"),
        (hire_cost_row, "Hiring Costs"),
        (total_exp_row, "Total Expenses"),
        (profit_row, "Net Profit / (Burn)"),
        (cash_row, "Cash Balance"),
    ]
    for r, label in row_labels:
        ws.cell(row=r, column=1, value=label)
        _style_cell(ws.cell(row=r, column=1), is_label=True)

    pricing_sheet_name = "'Pricing Scenarios'"

    for m in range(1, 37):
        col = m + 1
        cl = get_column_letter(col)

        # Revenue: months 1-12 reference Pricing sheet Total MRR, months 13+ extrapolate
        if m <= 12:
            pricing_col = get_column_letter(m + 1)
            rev_formula = f"={pricing_sheet_name}!{pricing_col}{pricing_total_mrr_row}"
        else:
            prev_cl = get_column_letter(col - 1)
            rev_formula = f"={prev_cl}{rev_row}*1.05"
        ws.cell(row=rev_row, column=col, value=rev_formula)
        _style_cell(ws.cell(row=rev_row, column=col), fmt=CURRENCY_FMT)

        # COGS = Revenue * COGS%
        cogs_formula = f"={cl}{rev_row}*$B$3"
        ws.cell(row=cogs_row, column=col, value=cogs_formula)
        _style_cell(ws.cell(row=cogs_row, column=col), fmt=CURRENCY_FMT)

        # Gross Margin = Revenue - COGS
        gross_formula = f"={cl}{rev_row}-{cl}{cogs_row}"
        ws.cell(row=gross_row, column=col, value=gross_formula)
        _style_cell(ws.cell(row=gross_row, column=col), fmt=CURRENCY_FMT)

        # Operating Expenses (with quarterly growth)
        if len(exp_labels) > 0:
            base_sum = "+".join([f"$B${exp_start + i}" for i in range(len(exp_labels))])
            quarter = f"INT(({m}-1)/3)"
            opex_formula = f"=({base_sum})*(1+$B$4)^{quarter}"
        else:
            opex_formula = "=0"
        ws.cell(row=opex_row, column=col, value=opex_formula)
        _style_cell(ws.cell(row=opex_row, column=col), fmt=CURRENCY_FMT)

        # Hiring Costs: SUM of IF(month >= start_month, salary, 0)
        if hires:
            hire_parts = []
            for i in range(len(hires)):
                salary_cell = f"$B${hire_data_start + 1 + i}"
                start_cell = f"$C${hire_data_start + 1 + i}"
                hire_parts.append(f"IF({m}>={start_cell},{salary_cell},0)")
            hire_formula = "=" + "+".join(hire_parts)
        else:
            hire_formula = "=0"
        ws.cell(row=hire_cost_row, column=col, value=hire_formula)
        _style_cell(ws.cell(row=hire_cost_row, column=col), fmt=CURRENCY_FMT)

        # Total Expenses
        total_exp_formula = f"={cl}{opex_row}+{cl}{hire_cost_row}"
        ws.cell(row=total_exp_row, column=col, value=total_exp_formula)
        _style_cell(ws.cell(row=total_exp_row, column=col), fmt=CURRENCY_FMT)

        # Net Profit / Burn
        profit_formula = f"={cl}{gross_row}-{cl}{total_exp_row}"
        ws.cell(row=profit_row, column=col, value=profit_formula)
        _style_cell(ws.cell(row=profit_row, column=col), fmt=CURRENCY_FMT)

        # Cash Balance
        if m == 1:
            cash_formula = f"=$B$5+{cl}{profit_row}"
        else:
            prev_cl = get_column_letter(col - 1)
            cash_formula = f"={prev_cl}{cash_row}+{cl}{profit_row}"
        ws.cell(row=cash_row, column=col, value=cash_formula)
        _style_cell(ws.cell(row=cash_row, column=col), fmt=CURRENCY_FMT)

    # Highlight key rows
    for r in [gross_row, profit_row, cash_row]:
        ws.cell(row=r, column=1).font = Font(name="Calibri", size=10, bold=True, color="1F4E79")

    # ── Line Chart ──
    chart = LineChart()
    chart.title = "Revenue, Profit & Cash (36 Months)"
    chart.style = 10
    chart.width = 36
    chart.height = 16
    chart.x_axis.title = "Month"
    chart.y_axis.title = currency

    from openpyxl.chart.series import SeriesLabel
    data_ref_rev = Reference(ws, min_col=2, max_col=37, min_row=rev_row, max_row=rev_row)
    data_ref_profit = Reference(ws, min_col=2, max_col=37, min_row=profit_row, max_row=profit_row)
    data_ref_cash = Reference(ws, min_col=2, max_col=37, min_row=cash_row, max_row=cash_row)
    cats = Reference(ws, min_col=2, max_col=37, min_row=header_row, max_row=header_row)

    chart.add_data(data_ref_rev, titles_from_data=False)
    chart.add_data(data_ref_profit, titles_from_data=False)
    chart.add_data(data_ref_cash, titles_from_data=False)
    chart.set_categories(cats)

    chart.series[0].title = SeriesLabel(v="Revenue")
    chart.series[0].graphicalProperties.line.solidFill = "10B981"
    chart.series[1].title = SeriesLabel(v="Profit / Burn")
    chart.series[1].graphicalProperties.line.solidFill = "EF4444"
    chart.series[2].title = SeriesLabel(v="Cash Balance")
    chart.series[2].graphicalProperties.line.solidFill = "3B82F6"

    chart_row = cash_row + 2
    ws.add_chart(chart, f"B{chart_row}")

    # Column widths
    ws.column_dimensions['A'].width = 24
    for c in range(2, 38):
        ws.column_dimensions[get_column_letter(c)].width = 12


# ═══════════════════════════════════════════════
# Sheet 3: Cap Table
# ═══════════════════════════════════════════════

def build_cap_table_sheet(wb: Workbook, cap_data: dict):
    ws = wb.create_sheet("Cap Table", 2)
    ws.sheet_properties.tabColor = "F59E0B"

    funding = cap_data.get("funding_amount", 500000)
    pre_money = cap_data.get("pre_money_valuation", 2000000)
    option_pool = cap_data.get("option_pool_pct", 0.10)
    use_of_funds = cap_data.get("use_of_funds", [])
    currency = cap_data.get("currency", "USD")

    # ── Section: Inputs ──
    ws.merge_cells('A1:D1')
    ws['A1'] = "CAP TABLE & FUNDING STRATEGY"
    ws['A1'].font = Font(name="Calibri", size=13, bold=True, color="1F4E79")

    ws['A3'] = "EDITABLE INPUTS (yellow cells)"
    ws['A3'].font = Font(name="Calibri", size=11, bold=True, color="B45309")

    inputs = [
        (4, f"Funding Amount ({currency})", funding, CURRENCY_FMT),
        (5, f"Pre-Money Valuation ({currency})", pre_money, CURRENCY_FMT),
        (6, "Option Pool %", option_pool, PCT_FMT),
    ]
    for row, label, val, fmt in inputs:
        ws.cell(row=row, column=1, value=label)
        _style_cell(ws.cell(row=row, column=1), is_label=True)
        ws.cell(row=row, column=2, value=val)
        _style_cell(ws.cell(row=row, column=2), fmt=fmt, is_input=True)

    # ── Section: Calculated Values ──
    ws['A8'] = "CALCULATED EQUITY STRUCTURE"
    ws['A8'].font = Font(name="Calibri", size=11, bold=True, color="1F4E79")

    calc_rows = [
        (9, f"Post-Money Valuation ({currency})", "=B4+B5", CURRENCY_FMT),
        (10, "Investor Ownership %", "=B4/B9", PCT_FMT),
        (11, "Option Pool %", "=B6", PCT_FMT),
        (12, "Founder Ownership %", "=1-B10-B11", PCT_FMT),
    ]
    for row, label, formula, fmt in calc_rows:
        ws.cell(row=row, column=1, value=label)
        _style_cell(ws.cell(row=row, column=1), is_label=True)
        ws.cell(row=row, column=2, value=formula)
        _style_cell(ws.cell(row=row, column=2), fmt=fmt)

    # Highlight ownership rows
    for r in [10, 12]:
        ws.cell(row=r, column=2).font = Font(name="Calibri", size=11, bold=True, color="1F4E79")

    # ── Section: Use of Funds ──
    ws['A14'] = "USE OF FUNDS BREAKDOWN"
    ws['A14'].font = Font(name="Calibri", size=11, bold=True, color="1F4E79")

    uof_headers = ["Category", "Allocation %", f"Amount ({currency})"]
    for i, h in enumerate(uof_headers, 1):
        ws.cell(row=15, column=i, value=h)
    _style_header(ws, 15, 3)

    if not use_of_funds:
        use_of_funds = [
            {"category": "Engineering", "allocation_pct": 0.40},
            {"category": "Marketing", "allocation_pct": 0.25},
            {"category": "Operations", "allocation_pct": 0.15},
            {"category": "Reserve", "allocation_pct": 0.20},
        ]

    for i, item in enumerate(use_of_funds):
        r = 16 + i
        ws.cell(row=r, column=1, value=item.get("category", ""))
        _style_cell(ws.cell(row=r, column=1), is_label=True)

        ws.cell(row=r, column=2, value=item.get("allocation_pct", 0))
        _style_cell(ws.cell(row=r, column=2), fmt=PCT_FMT, is_input=True)

        # Amount = Funding * allocation_pct (formula)
        formula = f"=$B$4*B{r}"
        ws.cell(row=r, column=3, value=formula)
        _style_cell(ws.cell(row=r, column=3), fmt=CURRENCY_FMT)

    # Total row
    total_r = 16 + len(use_of_funds)
    ws.cell(row=total_r, column=1, value="TOTAL")
    ws.cell(row=total_r, column=1).font = Font(name="Calibri", size=11, bold=True)
    sum_range_pct = f"B16:B{total_r - 1}"
    sum_range_amt = f"C16:C{total_r - 1}"
    ws.cell(row=total_r, column=2, value=f"=SUM({sum_range_pct})")
    _style_cell(ws.cell(row=total_r, column=2), fmt=PCT_FMT)
    ws.cell(row=total_r, column=3, value=f"=SUM({sum_range_amt})")
    _style_cell(ws.cell(row=total_r, column=3), fmt=CURRENCY_FMT)

    # Column widths
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 20


# ═══════════════════════════════════════════════
# Main Orchestrator
# ═══════════════════════════════════════════════

def generate_financial_model(plan_id: str, pricing_data: dict, projection_data: dict, cap_data: dict) -> str:
    """
    Generate the Financial_Model.xlsx workbook with all 3 sheets.
    Returns the file path on success, None on failure.
    """
    try:
        os.makedirs("./data/financials", exist_ok=True)
        file_path = f"./data/financials/{plan_id}.xlsx"

        wb = Workbook()
        # Remove default sheet
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]

        num_tiers = len(pricing_data.get("tiers", []))
        if num_tiers == 0:
            num_tiers = 2

        total_mrr_row = build_pricing_sheet(wb, pricing_data)
        build_projection_sheet(wb, projection_data, total_mrr_row, num_tiers)
        build_cap_table_sheet(wb, cap_data)

        wb.save(file_path)
        print(f"[Financial] Excel saved: {file_path}")
        return file_path

    except Exception as e:
        print(f"[Financial] Generation failed: {e}")
        import traceback
        traceback.print_exc()
        return None
